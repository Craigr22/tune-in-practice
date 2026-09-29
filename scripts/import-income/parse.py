#!/usr/bin/env python3
"""
Read a year's income workbook and plan an import into the portal.

    python3 scripts/import-income/parse.py "26-27 Income BAM (1).xlsx"

Writes to scripts/import-income/out/ (git-ignored — it is student data):

    plan.json      every row to insert, with a stable id each
    review.csv     one line per student, with everything that was guessed
    reconcile.txt  imported payments against each sheet's own monthly totals

Nothing touches the database. The plan is applied separately, after review.

Ids are uuid5 of the source — the sheet and the name — so running this again
on the same workbook produces the same ids. Applying twice inserts nothing new,
and undoing removes exactly what was added.

How the workbook is read, sheet by sheet:

  * A row whose name is a weekday and a time ("Thursday 8pm Guitar") opens a
    class. Students below it are in that class until the next one opens, or
    four blank rows go by.
  * A student with a monthly fee in column C is current. One without is from
    a class that has ended.
  * "students N" or "total" in the name column starts the list of past
    students; everyone below it is past, fee or not.
  * The month columns are payments. A negative is a refund.
  * A row with no name but a payment belongs to the student above it; a month
    cell holding a name ("shrutika") means the payment beside it is theirs.
"""

import datetime as dt
import json
import re
import sys
import uuid
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

TODAY = dt.date(2026, 9, 29)
FY_START_YEAR = 2026  # FY 2026-27: April 2026 to March 2027
MONTHS = ["apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar"]
NS = uuid.UUID("6f1c2a9e-0b1d-4e7a-9d3c-5a2e8b7c4d10")  # this importer's namespace

TEACHER_SHEETS = ["Brendan", "David", "Gaurangi", "Jason", "Garth", "Anurag", "Chris", "Jay", "Rushikesh", "Sachit"]
# Sheets that were hidden in the workbook: teachers no longer here.
FORMER_TEACHERS = {"Rushikesh", "Sachit"}

# What already exists in the portal, so it's matched rather than duplicated.
EXISTING = {
    "instruments": {
        "ukulele": "f5b4751a-9267-418c-9109-7bf6b4d879ab",
        "guitar": "b782abf7-5ef2-43b9-88f9-e09f0855d323",
        "violin": "ff551d6a-1a67-45d3-b545-2892c47168ad",
    },
    "locations": {
        "Andheri Room 1": "a27a6371-bef4-442c-9f27-ebafffc88ef2",
        "Andheri Room 2": "540f6034-dac8-4a39-a433-f2c63704a46d",
    },
    "teachers": {"Jason": "68f6b088-280b-4d49-b2fe-725739caa1f3"},
    # (teacher sheet, day, "HH:MM") -> existing batch
    "batches": {("Garth", 0, "15:00"): ("168a8bde-d5c9-40f0-8d22-ef9f603bcd17", "BAMUK01", "Jason")},
    "students": {
        "amit acharekar": "ec89ef20-b358-41f6-84dc-f4cdb568417a",
        "elroy rodrigues": "d5c67068-fa7a-4090-9447-7b8cec0606da",
        "errol ammana": "1a5087c1-6728-4418-acbe-2eb352b82ce2",
        "hetvi mathuria": "07fdbf39-fd4a-4a09-b61f-08071ad8eefe",
        "payal malviya": "15843720-fb24-490d-9f7b-6e9d59bc718a",
    },
}
DEFAULT_LOCATION = "Andheri Room 1"
DEFAULT_CAPACITY = 6

DAYS = {"sun": 0, "mon": 1, "tue": 2, "wed": 3, "thu": 4, "fri": 5, "sat": 6}
NOT_NAMES = {"students", "total", "paid but not started", "-", "we cancelled", "no one showed",
             "subtitue teacher", "substitute teacher", "class as usual", "1 on 1"}
SKIP_PEOPLE = {"craig rocha"}  # the owner, listed as a joke ("Bhai hai apna", expires "Infinity")
INSTRUMENT_WORDS = [("early music", "early music"), ("guitar", "guitar"), ("ukulele", "ukulele"),
                    ("violin", "violin"), ("piano", "piano")]


def sid(*parts):
    return str(uuid.uuid5(NS, "|".join(str(p) for p in parts)))


def norm(name):
    return re.sub(r"[^a-z]+", " ", str(name).lower()).strip()


def edit_distance(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


# ---------------------------------------------------------------- dates

MONTH_WORDS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def text_date_parts(s):
    """'6th Sep' / '23rd April' / '4th Oct' -> (day, month), else None."""
    m = re.search(r"(\d{1,2})\s*(?:st|nd|rd|th)?\s+([a-z]{3,})", str(s).lower())
    if not m:
        return None
    month = MONTH_WORDS.get(m.group(2)[:3])
    return (int(m.group(1)), month) if month else None


def safe_date(y, m, d):
    try:
        return dt.date(y, m, d)
    except ValueError:
        return None


def parse_start(v):
    """A start date: the latest reading that isn't in the future."""
    if isinstance(v, dt.datetime):
        return v.date(), "sheet"
    parts = text_date_parts(v) if isinstance(v, str) else None
    if not parts:
        return None, None
    options = [safe_date(y, parts[1], parts[0]) for y in (TODAY.year, TODAY.year - 1)]
    options = [d for d in options if d and d <= TODAY]
    return (max(options), "year inferred") if options else (None, None)


def parse_expiry(v, start):
    """An expiry date: the first reading on or after the start, else the one nearest today."""
    if isinstance(v, dt.datetime):
        return v.date(), "sheet"
    parts = text_date_parts(v) if isinstance(v, str) else None
    if not parts:
        return None, None
    options = [d for d in (safe_date(y, parts[1], parts[0]) for y in (TODAY.year - 1, TODAY.year, TODAY.year + 1)) if d]
    if start:
        after = [d for d in options if d >= start]
        if after:
            return min(after), "year inferred"
    return min(options, key=lambda d: abs((d - TODAY).days)), "year inferred"


def month_date(i):
    """First of the i-th month of the financial year (0 = April)."""
    month = (3 + i) % 12 + 1
    year = FY_START_YEAR if i < 9 else FY_START_YEAR + 1
    return dt.date(year, month, 1)


# ---------------------------------------------------------------- headings

def parse_heading(text):
    """'Monday Ukulele 7pm' -> (day, 'HH:MM', instrument or None), else None."""
    t = str(text).lower()
    day = next((d for w, d in DAYS.items() if re.match(rf"\s*{w}", t)), None)
    if day is None:
        return None
    tm = re.search(r"(\d{1,2})(?::(\d{2}))?\s*(am|pm)", t)
    if not tm:
        return None
    hour = int(tm.group(1)) % 12 + (12 if tm.group(3) == "pm" else 0)
    instrument = next((name for word, name in INSTRUMENT_WORDS if word in t), None)
    return day, f"{hour:02d}:{int(tm.group(2) or 0):02d}", instrument


def is_one_on_one_heading(text):
    return bool(re.search(r"1\s*on\s*1", str(text).lower()))


def split_name(raw):
    """'Aarav.Prashant Pai' -> ('Aarav', 'Prashant Pai', note). Nothing is invented."""
    s = str(raw).strip()
    note = None
    paren = re.search(r"\((.*)$", s)
    if paren:
        note = paren.group(1).rstrip(")").strip() or None
        s = s[: paren.start()].strip()
    parts = [p.strip() for p in re.split(r"\.", s) if p.strip()]
    name = parts[0] if parts else s
    parent = " ".join(parts[1:]) or None
    if name.islower():
        name = name.title()
    if parent and parent.islower():
        parent = parent.title()
    return name, parent, note


# ---------------------------------------------------------------- read

def read_sheet(ws, sheet):
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    april = next(i for i, c in enumerate(header) if isinstance(c, str) and c.strip().lower().startswith("apr"))
    summary = rows[1]
    sheet_months = [summary[april + k] if april + k < len(summary) else None for k in range(12)]
    sheet_months = [int(v) if isinstance(v, (int, float)) else 0 for v in sheet_months]

    students, headings = [], []
    current = None          # the class students are being read into
    blank_run = 0
    past = False            # below "students N" / "total"
    last_student = None

    for r_i, row in enumerate(rows[2:], start=3):
        if not any(c not in (None, "") for c in row):
            blank_run += 1
            if blank_run >= 4:
                current = None
            continue
        blank_run = 0

        a, b, c = row[0], row[1], row[2]
        d = row[3] if len(row) > 3 else None
        e = row[4] if april > 4 and len(row) > 4 else None
        cells = [row[april + k] if april + k < len(row) else None for k in range(12)]
        pays = {k: v for k, v in enumerate(cells) if isinstance(v, (int, float)) and v != 0}
        b_text = str(b).strip() if isinstance(b, str) else ""

        # A month cell holding a name: the payment beside it is theirs.
        named = [(k, str(v).strip()) for k, v in enumerate(cells) if isinstance(v, str) and v.strip()]
        if not b_text and named:
            for k, who in named:
                if k + 1 < 12 and isinstance(cells[k + 1], (int, float)) and cells[k + 1]:
                    students.append({"attach_to": who, "month": k + 1, "amount": cells[k + 1],
                                     "row": r_i, "sheet": sheet})
                    pays.pop(k + 1, None)
            if not pays:
                continue

        if b_text.lower() in ("students", "total"):
            past = True
            current = None
            continue

        heading = parse_heading(b_text) if b_text else None
        if heading and not isinstance(c, (int, float)) and not pays:
            current = {"key": (sheet, heading[0], heading[1]), "day": heading[0], "time": heading[1],
                       "instrument": heading[2], "text": b_text, "started": a, "row": r_i, "past": past}
            headings.append(current)
            continue
        if b_text and is_one_on_one_heading(b_text) and not isinstance(c, (int, float)):
            current = {"key": (sheet, "1on1"), "one_on_one": True, "text": b_text, "row": r_i,
                       "past": past, "instrument": None}
            headings.append(current)
            continue

        note_row = (not b_text or isinstance(b, (dt.datetime, dt.date))
                    or re.fullmatch(r"[\d/ ,.\-]+", b_text) or b_text.lower() in NOT_NAMES)
        if note_row:
            # A payment on a note row belongs to the student just above it.
            if pays and last_student:
                for k, amt in pays.items():
                    last_student["payments"].append({"month": k, "amount": amt, "row": r_i, "from_note": True})
            continue

        if norm(b_text) in SKIP_PEOPLE:
            continue

        name, parent, note = split_name(b_text)
        fee = c if isinstance(c, (int, float)) and c > 0 else None
        start, start_how = parse_start(a)
        expiry, expiry_how = parse_expiry(e, start) if e not in (None, "") else (None, None)
        student = {
            "sheet": sheet, "row": r_i, "raw": b_text, "name": name, "parent": parent, "note": note,
            "fee": fee, "missed": d if isinstance(d, (int, float)) else None,
            "start": start, "start_how": start_how, "expiry": expiry, "expiry_how": expiry_how,
            "a_text": a if isinstance(a, str) else None,
            "active": bool(fee) and not past,
            # Current, fee paying, and under no class heading: one of the teacher's
            # 1-on-1s (Anurag and Chris have no headings at all).
            "klass": None if past else (current or ({"key": (sheet, "1on1"), "one_on_one": True,
                                                     "text": "1-on-1", "row": r_i, "past": False,
                                                     "instrument": None} if fee else None)),
            "payments": [{"month": k, "amount": v, "row": r_i} for k, v in pays.items()],
        }
        students.append(student)
        last_student = student

    # Resolve the payments attached by name.
    real = [s for s in students if "attach_to" not in s]
    for tag in [s for s in students if "attach_to" in s]:
        hint = norm(tag["attach_to"])
        match = [s for s in real if hint and (hint in norm(s["raw"]))]
        target = match[-1] if match else None
        if target:
            target["payments"].append({"month": tag["month"], "amount": tag["amount"], "row": tag["row"],
                                       "by_name": tag["attach_to"]})
        else:
            print(f"  ! {tag['sheet']} row {tag['row']}: payment named '{tag['attach_to']}' matches no student")
    return real, headings, sheet_months


# ---------------------------------------------------------------- plan

def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = Path(__file__).parent / "out"
    out.mkdir(exist_ok=True)

    # Workshop sign-ups name an instrument; where one is also a course student,
    # it's evidence for what their class teaches.
    workshop = {}
    for row in list(wb["Workshop"].iter_rows(values_only=True))[2:]:
        if row[0] and row[1] and isinstance(row[1], str):
            workshop[norm(row[0])] = row[1].strip().lower().replace("early", "early music")
    workshop_total = sum(v for v in list(wb["Workshop"].iter_rows(values_only=True))[1][5:17]
                         if isinstance(v, (int, float)))

    all_students, all_headings, sheet_totals = [], [], {}
    for sheet in TEACHER_SHEETS:
        s, h, m = read_sheet(wb[sheet], sheet)
        all_students += s
        all_headings += h
        sheet_totals[sheet] = m

    plan = {"instruments": [], "teachers": [], "batches": [], "students": [], "enrollments": [],
            "payments": [], "student_updates": []}
    flags = defaultdict(list)  # student key -> reasons for a second look

    # ---- instruments per class
    classes = {}
    for s in all_students:
        if s["active"] and s["klass"]:
            k = s["klass"]["key"]
            classes.setdefault(k, {"heading": s["klass"], "students": []})["students"].append(s)
    for k, c in classes.items():
        h = c["heading"]
        if h.get("instrument"):
            c["instrument"], c["how"] = h["instrument"], "named in the class heading"
            continue
        seen = Counter(workshop[norm(x["raw"])] for x in c["students"] if norm(x["raw"]) in workshop)
        seen += Counter(workshop[norm(x["name"])] for x in c["students"] if norm(x["name"]) in workshop)
        if seen:
            c["instrument"], c["how"] = seen.most_common(1)[0][0], "a student's workshop sign-up"
    by_teacher = defaultdict(Counter)
    for k, c in classes.items():
        if c.get("instrument"):
            by_teacher[k[0]][c["instrument"]] += 1
    teacher_default = {t: cnt.most_common(1)[0][0] for t, cnt in by_teacher.items()}
    teacher_default.setdefault("Jason", "ukulele")
    for k, c in classes.items():
        if not c.get("instrument"):
            guess = teacher_default.get(k[0])
            c["instrument"] = guess or "unassigned"
            c["how"] = "the teacher's other classes" if guess else "unknown — set it"

    instrument_ids = dict(EXISTING["instruments"])
    for name in sorted({c["instrument"] for c in classes.values()} - set(instrument_ids)):
        iid = sid("instrument", name)
        instrument_ids[name] = iid
        plan["instruments"].append({"id": iid, "name": name, "is_active": False})

    # ---- teachers
    teacher_ids = {}
    for t in TEACHER_SHEETS:
        if t in EXISTING["teachers"]:
            teacher_ids[t] = EXISTING["teachers"][t]
            continue
        teacher_ids[t] = sid("teacher", t)
        taught = sorted({instrument_ids[c["instrument"]] for k, c in classes.items() if k[0] == t})
        plan["teachers"].append({
            "id": teacher_ids[t], "name": t, "is_active": t not in FORMER_TEACHERS,
            "payment_type": "per_session", "rate": 0, "instruments": taught,
        })

    # ---- batches
    batch_ids = {}
    for k, c in sorted(classes.items(), key=lambda kv: (kv[0][0], str(kv[0][1:]))):
        h = c["heading"]
        teacher = k[0]
        if h.get("one_on_one"):
            day, time = 0, "00:00"
        else:
            day, time = h["day"], h["time"]
        existing = EXISTING["batches"].get((teacher, day, time))
        if existing:
            batch_ids[k] = existing[0]
            c["existing"] = existing
            continue
        bid = sid("batch", *k)
        batch_ids[k] = bid
        plan["batches"].append({
            "id": bid, "teacher_id": teacher_ids[teacher], "instrument_id": instrument_ids[c["instrument"]],
            "location_id": EXISTING["locations"][DEFAULT_LOCATION], "day_of_week": day,
            "start_time": time, "duration_min": 60,
            "max_students": 1 if h.get("one_on_one") else max(DEFAULT_CAPACITY, len(c["students"])),
            "is_active": teacher not in FORMER_TEACHERS,
            # no semester_start: a class with one generates twelve weeks of sessions on insert
            "_label": f"{teacher} · {h['text']}", "_instrument_how": c["how"], "_students": len(c["students"]),
        })

    # ---- students: one record per person, across sheets
    people = {}
    for s in all_students:
        people.setdefault(norm(s["name"] if not s["parent"] else s["raw"]), []).append(s)
    existing_keys = EXISTING["students"]
    outcome = {}  # person key -> what was planned for them, for the review sheet

    for key, entries in people.items():
        primary = next((x for x in entries if x["active"]), entries[0])
        match = existing_keys.get(norm(primary["name"])) or existing_keys.get(key)
        if not match:
            near = [(edit_distance(norm(primary["name"]), n), n) for n in existing_keys]
            near = [x for x in near if x[0] <= 2]
            if near:
                match = existing_keys[min(near)[1]]
                flags[key].append(f"matched to existing portal student '{min(near)[1]}' (spelling differs)")
        student_id = match or sid("student", key)

        active = any(x["active"] for x in entries)
        fee = primary["fee"]
        pays = [p for x in entries for p in x["payments"]]
        # The latest month anyone paid in, taken whole: a quarter is often paid in
        # two parts in the same month (₹4,000 then ₹8,000), and neither part alone
        # says what cycle they're on.
        latest_month = max((p["month"] for p in pays if p["amount"] > 0), default=None)
        latest_pays = [p for p in pays if p["month"] == latest_month and p["amount"] > 0]
        latest = {"amount": sum(p["amount"] for p in latest_pays), "month": latest_month} if latest_pays else None
        cycle, fee_amount = "monthly", fee or 0
        if fee and latest:
            ratio = latest["amount"] / fee
            if 2.5 <= ratio <= 3.5:
                cycle, fee_amount = "quarterly", fee * 3
            elif not 0.8 <= ratio <= 1.2:
                flags[key].append(f"last payment ₹{latest['amount']:,.0f} doesn't fit a ₹{fee:,.0f} monthly fee")
        joined = primary["start"] or (month_date(min(p["month"] for p in pays)) if pays else dt.date(FY_START_YEAR, 4, 1))
        if not primary["start"] and not pays and active:
            flags[key].append("joined date unknown — set to 1 Apr 2026")

        if len(entries) > 1:
            flags[key].append("appears on " + ", ".join(f"{x['sheet']} row {x['row']}" for x in entries))
        if primary["note"]:
            flags[key].append(f"note from the sheet: ({primary['note']})")
        # A guessed year only matters where it sets a current student's period.
        # "4th Oct" for a current student is this October; only an odd result needs a look.
        if primary["active"] and primary["expiry_how"] == "year inferred" and primary["expiry"] and \
                not -45 <= (primary["expiry"] - TODAY).days <= 150:
            flags[key].append(f"expiry '{primary['expiry']}' had no year in the sheet; inferred year looks odd")

        if match:
            plan["student_updates"].append({"id": student_id, "fee_amount": fee_amount, "fee_cycle": cycle,
                                            "_name": primary["name"], "_only_if_unset": True})
        else:
            plan["students"].append({
                "id": student_id, "name": primary["name"], "parent_name": primary["parent"],
                "fee_amount": fee_amount, "fee_cycle": cycle, "is_active": active,
                "joined_on": joined.isoformat(),
            })

        # enrolment: current students only, in their class
        batch_id = None
        if primary["active"] and primary["klass"]:
            batch_id = batch_ids.get(primary["klass"]["key"])
            if batch_id and not (match and batch_id == EXISTING["batches"].get((primary["sheet"], 0, "15:00"), [None])[0]):
                plan["enrollments"].append({
                    "id": sid("enrollment", student_id, batch_id), "student_id": student_id,
                    "batch_id": batch_id, "status": "active", "enrolled_on": joined.isoformat(),
                })
        elif primary["active"]:
            flags[key].append("current, but not under a class heading — not enrolled anywhere")

        outcome[key] = {"id": student_id, "batch_id": batch_id, "primary": primary, "entries": entries,
                        "fee_amount": fee_amount, "cycle": cycle, "existing": bool(match)}

        # payments
        for p in pays:
            when = month_date(p["month"])
            is_latest = latest is not None and p in latest_pays and primary["active"]
            period_start = period_end = None
            if is_latest and primary["expiry"]:
                period_start = primary["start"] or when
                period_end = primary["expiry"]
                if period_end < period_start:
                    flags[key].append(f"expiry {period_end} is before start {period_start}; period left blank")
                    period_start = period_end = None
            source = f"{primary['sheet'] if p.get('row') == primary['row'] else next((x['sheet'] for x in entries if x['row'] == p['row']), primary['sheet'])}"
            note = f"Imported from 26-27 Income BAM.xlsx · {source} row {p['row']} · {MONTHS[p['month']].title()}"
            if p.get("from_note"):
                note += " · on a note row, assumed to be this student's"
                flags[key].append(f"₹{p['amount']:,.0f} in {MONTHS[p['month']].title()} was on a note row below them")
            if p.get("by_name"):
                note += f" · marked '{p['by_name']}'"
            if p["amount"] < 0:
                note += " · refund"
            plan["payments"].append({
                "id": sid("payment", student_id, p["row"], p["month"], p["amount"]),
                "student_id": student_id, "batch_id": batch_id, "amount": p["amount"],
                "paid_on": when.isoformat(),
                "period_start": period_start.isoformat() if period_start else None,
                "period_end": period_end.isoformat() if period_end else None,
                "method": None, "status": "paid", "notes": note,
                "_sheet": source,
            })

    (out / "plan.json").write_text(json.dumps(plan, indent=1, default=str))

    # ---- review sheet
    import csv
    labels = {b["id"]: b["_label"] for b in plan["batches"]}
    for (teacher, day, time), (bid, code, _) in EXISTING["batches"].items():
        labels[bid] = f"{code} (already in the portal)"
    with open(out / "review.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Teacher", "Class", "Student", "Parent", "Status", "Fee", "Cycle", "Current period",
                    "Payments this year", "Check", "Source"])
        for key, o in sorted(outcome.items(), key=lambda kv: (kv[1]["primary"]["sheet"], kv[1]["primary"]["row"])):
            pr = o["primary"]
            theirs = sorted((p for p in plan["payments"] if p["student_id"] == o["id"]), key=lambda p: p["paid_on"])
            period = next((f"{p['period_start']} to {p['period_end']}" for p in theirs if p["period_end"]), "")
            w.writerow([
                pr["sheet"], labels.get(o["batch_id"], ""), pr["name"], pr["parent"] or "",
                ("current" if pr["active"] else "past") + (" · existing" if o["existing"] else ""),
                o["fee_amount"], o["cycle"], period,
                ", ".join(f"{p['paid_on'][:7]} {p['amount']:,.0f}" for p in theirs),
                " | ".join(flags.get(key, [])),
                "; ".join(f"{x['sheet']} row {x['row']}" for x in o["entries"]),
            ])

    # ---- reconciliation
    lines = []
    by_sheet = defaultdict(lambda: [0] * 12)
    for p in plan["payments"]:
        k = MONTHS.index(dt.date.fromisoformat(p["paid_on"]).strftime("%b").lower())
        by_sheet[p["_sheet"]][k] += p["amount"]
    ok = True
    formulas = openpyxl.load_workbook(path)  # formulas rather than values
    for sheet in TEACHER_SHEETS:
        f = str(formulas[sheet].cell(2, 6).value or "")
        m = re.search(r"SUM\([A-Z]+(\d+):", f)
        differs = [int(x) for x in by_sheet[sheet]] != sheet_totals[sheet]
        if m and int(m.group(1)) > 3 and differs:
            lines.append(f"{sheet:<10} NOTE: the sheet's own totals are {f} — rows 3–{int(m.group(1)) - 1} "
                         f"are left out of them, so the sheet under-reports. The import counts every row.")
        got, want = [int(x) for x in by_sheet[sheet]], sheet_totals[sheet]
        diff = [(MONTHS[i], got[i], want[i]) for i in range(12) if got[i] != want[i]]
        status = "matches" if not diff else "DIFFERS"
        ok &= not diff
        lines.append(f"{sheet:<10} imported ₹{sum(got):>9,}  sheet ₹{sum(want):>9,}  {status}")
        for m, g, s in diff:
            lines.append(f"            {m}: imported {g:,} vs sheet {s:,}")
    lines.append(f"\nWorkshop sign-ups (not imported — one-off, not course fees): ₹{workshop_total:,.0f}")
    (out / "reconcile.txt").write_text("\n".join(lines) + "\n")

    # ---- summary
    print("\n".join(lines))
    print(f"\nPlan: {len(plan['teachers'])} teachers, {len(plan['instruments'])} instruments, "
          f"{len(plan['batches'])} classes, {len(plan['students'])} new students "
          f"({sum(s['is_active'] for s in plan['students'])} current), "
          f"{len(plan['student_updates'])} existing students updated, "
          f"{len(plan['enrollments'])} enrolments, {len(plan['payments'])} payments "
          f"(₹{sum(p['amount'] for p in plan['payments']):,.0f}).")
    print(f"Flagged for a second look: {sum(1 for v in flags.values() if v)} students.")
    print("\nClasses:")
    for b in plan["batches"]:
        print(f"  {b['_label']:<45} {b['_students']:>2} students · instrument from {b['_instrument_how']}")
    for k, c in classes.items():
        if c.get("existing"):
            print(f"  {k[0]} · {c['heading']['text']:<32} -> existing {c['existing'][1]} (portal teacher: {c['existing'][2]})")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "26-27 Income BAM (1).xlsx"))
